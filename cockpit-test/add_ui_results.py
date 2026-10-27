# -*- coding: utf-8 -*-
"""把 ui-click-test 的实测结果写入用例工作簿，新增 12_网站点击实测 页。"""

import json
import os

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, "智能座舱域控测试用例集_V1.0.xlsx")
XLSX_OUT = os.path.join(BASE, "智能座舱域控测试用例集_V1.1.xlsx")
RESULTS = os.path.join(BASE, "report", "results.json")

PATHS = {
    "U-001": "打开 index.html → 等待扫描完成 → 读顶部扫描条",
    "U-002": "依次点击左侧 CDC / MDC / VDC 三个域按钮",
    "U-003": "每个域点击『诊断总览』+ 逐个点击全部模块（23 个）",
    "U-004": "进入 SOC 模块 → 点击『严重』筛选 → 点击『异常』恢复",
    "U-005": "点击『导入日志』→ 粘贴 5 行自定义日志 → 解析并刷新 → 回 SOC 模块查看时间轴",
    "U-006": "点击右上『一键导出维修报告』→ 等待下载 → 检查文件与内嵌图表",
    "U-007": "点击右上主题切换按钮两次",
    "U-008": "Emulation 依次设 1920/1440/1024/390 宽度 → 检查 scrollWidth",
    "U-009": "全程监听 Runtime.consoleAPICalled / Log.entryAdded",
}

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
PASS_FILL = PatternFill("solid", fgColor="E2EFDA")
FAIL_FILL = PatternFill("solid", fgColor="FCE4E4")
WRAP = Alignment(wrap_text=True, vertical="top")


def main():
    with open(RESULTS, encoding="utf-8") as f:
        results = json.load(f)

    wb = load_workbook(XLSX)
    if "12_网站点击实测" in wb.sheetnames:
        del wb["12_网站点击实测"]
    ws = wb.create_sheet("12_网站点击实测")

    widths = [9, 26, 46, 50, 9, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells("A1:F1")
    ws["A1"] = "诊断网站界面点击实测（Chrome 无头模式 · CDP 驱动 · 真实点击）"
    ws["A1"].font = Font(bold=True, size=13, color="1F3864")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:F2")
    ws["A2"] = ("执行方式：本地 http 服务 + Chrome headless + DevTools Protocol 真实派发 click；"
                "分辨率 1920×1080；日志来自内置 66 行演示数据与 5 行自定义导入。"
                "对应脚本：cockpit-test/ui-click-test.mjs")
    ws["A2"].alignment = WRAP
    ws["A2"].font = Font(size=9, color="595959")
    ws.row_dimensions[2].height = 34

    head = ["用例ID", "用例名称", "点击/操作路径", "实测结果（脚本断言输出）", "结论", "耗时(ms)"]
    for i, h in enumerate(head, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[3].height = 26

    r = 4
    for item in results:
        row = [item["id"], item["name"], PATHS.get(item["id"], ""), item["detail"],
               "Pass" if item["status"] == "Pass" else "Fail", item["ms"]]
        for i, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(size=10)
            c.alignment = WRAP
            c.border = BORDER
        ws.cell(row=r, column=5).fill = PASS_FILL if item["status"] == "Pass" else FAIL_FILL
        ws.cell(row=r, column=5).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=r, column=6).alignment = Alignment(horizontal="center", vertical="top")
        ws.row_dimensions[r].height = 46
        r += 1

    passed = sum(1 for x in results if x["status"] == "Pass")
    ws.cell(row=r, column=1, value="合计").font = Font(bold=True)
    ws.cell(row=r, column=2, value=f"{passed}/{len(results)} 通过").font = Font(bold=True)
    for c in range(1, 7):
        ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="DCE6F1")
        ws.cell(row=r, column=c).border = BORDER

    # 截图
    shots = [
        (os.path.join(BASE, "report", "shots", "01-overview-cdc.png"), r + 3, 1, "CDC 域诊断总览（1920×1080）"),
        (os.path.join(BASE, "report", "shots", "07-mobile-390.png"), r + 3, 2, "390px 窄屏"),
    ]
    for path, row, col, caption in shots:
        if not os.path.exists(path):
            continue
        img = XLImage(path)
        if col == 1:
            img.width, img.height = 860, int(860 * img.height / img.width)
        else:
            img.width, img.height = 300, int(300 * img.height / img.width)
        anchor = f"{get_column_letter(col)}{row}"
        ws.add_image(img, anchor)

    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(XLSX_OUT)
    print("updated:", XLSX_OUT)
    print("sheets:", wb.sheetnames)


if __name__ == "__main__":
    main()
